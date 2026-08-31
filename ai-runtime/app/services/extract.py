from __future__ import annotations

from pathlib import Path

from app.schemas import ExtractResponse

MAX_TEXT_CHARS = 500_000
MAX_PDF_PAGES = 200
MAX_XLSX_ROWS = 10_000
MAX_XLSX_SHEETS = 50

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".tif", ".tiff"}
TEXT_EXT = {".txt", ".csv", ".md", ".rtf"}


def extract_path(path: str, original_name: str) -> ExtractResponse:
    ext = Path(original_name or path).suffix.lower()
    if ext in IMAGE_EXT:
        return ExtractResponse(
            engine="none",
            text="",
            page_count=1,
            confidence=0.0,
            warnings=["ocr_required"],
        )
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".xlsx":
        return _extract_xlsx(path)
    if ext in TEXT_EXT:
        return _extract_text(path, ext)
    return ExtractResponse(
        engine="none",
        text="",
        page_count=0,
        confidence=0.0,
        warnings=["unsupported_type"],
    )


def _clip(text: str) -> str:
    if len(text) <= MAX_TEXT_CHARS:
        return text
    return text[:MAX_TEXT_CHARS]


def _extract_pdf(path: str) -> ExtractResponse:
    from pypdf import PdfReader

    warnings: list[str] = []
    reader = PdfReader(path, strict=False)
    pages = min(len(reader.pages), MAX_PDF_PAGES)
    if len(reader.pages) > MAX_PDF_PAGES:
        warnings.append("page_limit")
    parts: list[str] = []
    for i in range(pages):
        try:
            parts.append(reader.pages[i].extract_text() or "")
        except Exception:
            warnings.append("page_error")
    text = _clip("\n".join(parts).strip())
    confidence = 0.8 if text else 0.2
    if not text:
        warnings.append("empty_text")
    return ExtractResponse(engine="pypdf", text=text, page_count=pages, confidence=confidence, warnings=warnings)


def _extract_docx(path: str) -> ExtractResponse:
    from docx import Document

    warnings: list[str] = []
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text]
    text = _clip("\n".join(parts).strip())
    if not text:
        warnings.append("empty_text")
    return ExtractResponse(
        engine="python-docx",
        text=text,
        page_count=max(1, len(doc.paragraphs) // 40),
        confidence=0.85 if text else 0.2,
        warnings=warnings,
    )


def _extract_xlsx(path: str) -> ExtractResponse:
    from openpyxl import load_workbook

    warnings: list[str] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = wb.sheetnames[:MAX_XLSX_SHEETS]
    if len(wb.sheetnames) > MAX_XLSX_SHEETS:
        warnings.append("sheet_limit")
    parts: list[str] = []
    rows_total = 0
    for name in sheets:
        parts.append(f"=== {name} ===")
        ws = wb[name]
        for row in ws.iter_rows(values_only=True):
            rows_total += 1
            if rows_total > MAX_XLSX_ROWS:
                warnings.append("row_limit")
                break
            cells = ["" if c is None else str(c) for c in row]
            parts.append("\t".join(cells).rstrip())
        if rows_total > MAX_XLSX_ROWS:
            break
    wb.close()
    text = _clip("\n".join(parts).strip())
    if not text:
        warnings.append("empty_text")
    return ExtractResponse(
        engine="openpyxl",
        text=text,
        page_count=len(sheets),
        confidence=0.85 if text else 0.2,
        warnings=warnings,
    )


def _extract_text(path: str, ext: str) -> ExtractResponse:
    warnings: list[str] = []
    data = Path(path).read_bytes()[: MAX_TEXT_CHARS * 4]
    text = ""
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        warnings.append("decode_failed")
        return ExtractResponse(engine="text", text="", page_count=0, confidence=0.0, warnings=warnings)
    if ext == ".rtf" and not text.lstrip().startswith("{\\rtf"):
        warnings.append("rtf_unparsed")
    text = _clip(text)
    return ExtractResponse(engine="text", text=text, page_count=1, confidence=0.9, warnings=warnings)
